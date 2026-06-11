import sys
from pathlib import Path
from docx import Document

def dump_docx(path, out):
    p = Path(path)
    doc = Document(p)
    out.write(f"=== {p.name} ===\n")
    out.write(f"PARAGRAPHS: {len(doc.paragraphs)}\n")
    for i, para in enumerate(doc.paragraphs):
        style = para.style.name if para.style else ""
        text = para.text
        if text.strip() or style.startswith("Heading"):
            out.write(f"[P{i:04d}|{style}] {text}\n")
    out.write(f"\nTABLES: {len(doc.tables)}\n")
    for ti, t in enumerate(doc.tables):
        out.write(f"-- TABLE {ti} rows={len(t.rows)} cols={len(t.columns)} --\n")
        for ri, row in enumerate(t.rows):
            cells = [c.text.replace("\n", " | ") for c in row.cells]
            out.write(f"  R{ri}: " + " || ".join(cells) + "\n")
    out.write(f"\nSECTIONS: {len(doc.sections)}\n")
    for si, s in enumerate(doc.sections):
        out.write(f"  S{si}: page={s.page_width}x{s.page_height} margins L={s.left_margin} R={s.right_margin} T={s.top_margin} B={s.bottom_margin}\n")

if __name__ == "__main__":
    paths = sys.argv[1:]
    out = sys.stdout
    for p in paths:
        try:
            dump_docx(p, out)
        except Exception as e:
            out.write(f"!! Failed {p}: {e}\n")
